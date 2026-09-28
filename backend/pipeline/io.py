import csv
import sys
from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import load_workbook

from pipeline.config import RAW_DIR, RAW_FILES

ENCODINGS = ("utf-8-sig", "cp949", "euc-kr")


class ColumnError(Exception):
    pass


def _clean_header(fields: Iterable[str | None]) -> list[str]:
    return [(f or "").strip().replace("\ufeff", "") for f in fields]


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    last: Exception | None = None
    for enc in ENCODINGS:
        try:
            with path.open(encoding=enc, newline="") as f:
                reader = csv.reader(f)
                header = _clean_header(next(reader))
                rows: list[dict[str, str]] = []
                for lineno, row in enumerate(reader, start=2):
                    if not any(row):
                        continue
                    cells = [c.strip() for c in row]
                    while len(cells) > len(header) and cells[-1] == "":
                        cells.pop()
                    if len(cells) != len(header):
                        raise ColumnError(
                            f"{path.name} {lineno}행: 컬럼 {len(cells)}개, 헤더 {len(header)}개"
                        )
                    rows.append(dict(zip(header, cells)))
                return rows
        except UnicodeDecodeError as exc:
            last = exc
    raise UnicodeDecodeError("csv", b"", 0, 0, f"{path}: {last}")


def read_xlsx_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = _clean_header(next(rows))
    out = []
    for row in rows:
        if row is None or all(v in (None, "") for v in row):
            continue
        out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(header, row)})
    wb.close()
    return out


def pick(row: dict[str, str], aliases: Sequence[str], required: bool = True) -> str:
    for a in aliases:
        if a in row:
            return row[a]
    if required:
        raise ColumnError(f"컬럼 {list(aliases)} 없음. 실제 헤더: {list(row)}")
    return ""


def write_csv(path: Path, fieldnames: Sequence[str], rows: Iterable[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(fieldnames))
        writer.writeheader()
        for r in rows:
            writer.writerow(r)


def require_raw(key: str) -> Path:
    spec = RAW_FILES[key]
    path = RAW_DIR / spec.filename
    if not path.exists():
        print(f"원본 파일 없음: {path}", file=sys.stderr)
        print(f"  다운로드: {spec.url}", file=sys.stderr)
        print(f"  저장 위치: {path}", file=sys.stderr)
        sys.exit(2)
    return path
